"""Lectura de datos y cálculos de la Ficha territorial y sectorial (sin Streamlit).

Todo lo que no es presentación vive aquí, para que la app, el script de validación
(scripts/validar_datos.py) y las pruebas (tests/) usen exactamente el mismo código.
"""
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# Valores que se tratan como "sin dato" (ya normalizados con norm_text)
SIN_DATO = {"", "NAN", "NONE", "NO REPORTA", "SIN INFORMACION", "NO APLICA", "N A", "NA"}
# Destinos de EQUIVALENCIAS que no asignan el registro a ninguna ficha sectorial
DESTINOS_SIN_SECTOR = {"SIN SECTOR", "TERRITORIAL", "POR PROYECTO", ""}
# Datos personales que no se publican (Ley 1581 de 2012); se eliminan al cargar
COLUMNAS_NO_PUBLICAS = ["Nombre enlace cooperación", "Correo del enlace de Cooperación"]
# Campos de la información general que no se muestran en el registro del departamento (normalizados)
CAMPOS_OCULTOS_REGISTRO = {"DIVIPOLA", "PORCENTAJE DE AVANCE", "DEPT NORM", "SECTOR FICHA"}
AMBITO_NACIONAL = "Ámbito Nacional"

# Columnas que la app necesita de cada fuente (se comparan sin tildes ni mayúsculas)
COLUMNAS_AOD = ["CODIGO INTERVENCION", "NOMBRE INTERVENCION", "OBJETIVO GENERAL", "FECHA INICIAL", "FECHA FINAL",
                "DEPARTAMENTO", "MUNICIPIO", "VALOR APORTE (USD)", "NOMBRE ACTOR", "ODS", "ENCI PRIMER NIVEL",
                "SECTORES GOB"]
ESQUEMA = {
    "infogeneral": ["Divipola", "Departamento", "Capital", "Número de Municipios", "Población"],
    "aod": COLUMNAS_AOD,
    "aod_ant": COLUMNAS_AOD,
    "aod_sectores": COLUMNAS_AOD,
    "css": ["Código", "NOMBRE DE LA INICIATIVA", "OBJETIVO GENERAL/DESCRIPCION DE LA INICIATIVA", "PAIS SOCIO",
            "REGION", "VIA DE COOPERACION", "MODALIDAD", "FECHA DE APROBACION", "ESPACIO VINCULADO",
            "ENTIDAD(ES) NACIONAL(ES)", "ENTIDAD(ES) EXTRANJERA(S)"],
    "colcol": ["CODIGO", "NOMBRE DEL INTERCAMBIO", "OBJETIVO DEL INTERCAMBIO", "ETAPA", "LÍNEA TEMÁTICA",
               "AÑO DE REALIZACIÓN", "ENTIDAD SOCIA NACIONAL", "DEPARTAMENTOS PARTICIPANTES",
               "PRESUPUESTO ESTIMADO APC COLOMBIA"],
    "contrapartidas": ["Departamento", "AÑO", "Nombre del proyecto", "Implementador", "Financiador",
                       "Monto total", "Monto por APC"],
    "info_sectores": ["Nombre del sector", "Cabeza de sector"],
    "colcol_sectores": ["CODIGO", "SECTOR VINCULADO"],
}


# =============================================================
# Texto y formatos
# =============================================================
def norm_text(x):
    """Mayúsculas, sin tildes ni signos, espacios simples. Sirve para comparar nombres."""
    if x is None:
        return ""
    s = str(x).strip().upper()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^A-Z0-9\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def num_es(v, decimales=0):
    """Número con separadores colombianos: 1.234,5"""
    return f"{v:,.{decimales}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def format_usd(n):
    try:
        return "USD " + num_es(float(n))
    except (TypeError, ValueError):
        return ""


def format_cop(n):
    try:
        return "$ " + num_es(float(n))
    except (TypeError, ValueError):
        return ""


def valor_corto(v):
    """Valor abreviado para tarjetas y etiquetas: 1.772 M, 72,5 M, 350 mil."""
    v = float(v or 0)
    if v >= 1e8:
        return num_es(v / 1e6) + " M"
    if v >= 1e6:
        return num_es(v / 1e6, 1) + " M"
    if v >= 1e4:
        return num_es(v / 1e3) + " mil"
    return num_es(v)


def titulo_es(texto):
    """Mayúscula inicial en cada palabra, sin capitalizar artículos ni conectores."""
    t = "; ".join(p.strip() for p in str(texto).split(";")).title()
    for palabra in (" Y ", " E ", " De ", " Del ", " La ", " Las ", " Los ", " Para ", " En "):
        t = t.replace(palabra, palabra.lower())
    return t


def fecha_larga(f):
    f = pd.Timestamp(f)
    return f"{f.day} de {MESES[f.month - 1]} de {f.year}"


def variacion_txt(act, ant):
    """Variación porcentual frente al corte anterior."""
    if ant == 0:
        return "Nuevo" if act > 0 else "-"
    p = (act - ant) / ant * 100
    if abs(p) < 0.5:
        return "Sin cambio"
    flecha = "▲" if p >= 0 else "▼"
    return f"{flecha} {abs(p):.0f}%"


def get_col(row, *names):
    """Valor de la primera columna con dato entre varios nombres posibles."""
    for name in names:
        val = row.get(name, None)
        if val is not None and str(val).strip() not in ("", "None", "nan"):
            return val
    return ""


def buscar_col(df, *nombres):
    """Nombre real de la primera columna que coincida (sin importar tildes, mayúsculas ni espacios)."""
    reales = {norm_text(c): c for c in df.columns}
    for nombre in nombres:
        if norm_text(nombre) in reales:
            return reales[norm_text(nombre)]
    return None


# =============================================================
# Lectura de fuentes
# =============================================================
def read_named_table(wb, table_name):
    """Lee una tabla con nombre (Insertar > Tabla) de un libro ya abierto."""
    for ws in wb.worksheets:
        if table_name in ws.tables:
            min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
            data = [list(row) for row in ws.iter_rows(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)]
            return pd.DataFrame(data[1:], columns=data[0])
    raise KeyError(f"No se encontró la tabla '{table_name}' en el libro.")


def limpiar_texto(df):
    """Columnas de texto como str sin espacios sobrantes; valor de aporte numérico."""
    df = df.copy()
    for c in df.columns:
        if df[c].dtype == "object":
            df[c] = df[c].astype(str).str.strip()
    if "VALOR APORTE (USD)" in df.columns:
        df["VALOR APORTE (USD)"] = pd.to_numeric(df["VALOR APORTE (USD)"], errors="coerce").fillna(0)
    return df


def cargar_config(ruta):
    """Cortes, fuentes, sectores ocultos, equivalencias y registros desde configuracion.xlsx."""
    hojas = pd.read_excel(ruta, sheet_name=None)
    texto = lambda v: "" if pd.isna(v) else str(v).strip()
    cortes = {
        texto(r["Rol"]).lower(): {
            "etiqueta": texto(r["Etiqueta"]),
            "vigencia": pd.Timestamp(r["Fecha de vigencia"]),
            "extraccion": pd.Timestamp(r["Fecha de extracción"]),
            "hoja": texto(r["Hoja en Ficha sectores.xlsx"]),
        }
        for _, r in hojas["CORTES"].iterrows()
    }
    fuentes = {
        texto(r["Fuente"]): {"actualizacion": texto(r["Actualización"]), "responsable": texto(r["Responsable"])}
        for _, r in hojas["FUENTES"].iterrows()
    }
    ocultos = {
        norm_text(r["Sector (como en INFO GENERAL)"])
        for _, r in hojas["SECTORES"].iterrows() if norm_text(r["Mostrar en la ficha"]) == "NO"
    }
    equivalencias = {
        (texto(r["Fuente"]).upper(), norm_text(r["Valor en la fuente"])): norm_text(r["Sector en la ficha"])
        for _, r in hojas["EQUIVALENCIAS"].iterrows()
    }
    registros = {
        (texto(r["Fuente"]).upper(), texto(r["Código"])): norm_text(r["Sector en la ficha"])
        for _, r in hojas["REGISTROS"].iterrows() if texto(r["Sector en la ficha"])
    }
    return {"cortes": cortes, "fuentes": fuentes, "ocultos": ocultos,
            "equivalencias": equivalencias, "registros": registros}


def cargar_xlsm(ruta):
    """Información general, Sur-Sur, ColCol y Contrapartidas del libro Ficha_territorial.xlsm.

    Se abre una sola vez (abrirlo cuesta varios segundos). La AOD se lee de Ficha sectores.xlsx.
    """
    wb = load_workbook(ruta, data_only=True)
    info = read_named_table(wb, "infogeneral")
    info = info.drop(columns=[c for c in info.columns
                              if norm_text(c) in {norm_text(x) for x in COLUMNAS_NO_PUBLICAS}])
    filas_css = [list(r) for r in wb["CSS"].iter_rows(values_only=True)]
    contrapartidas = read_named_table(wb, "contrapartidas")
    contrapartidas.columns = [str(c).strip().strip("'") for c in contrapartidas.columns]
    return {
        "infogeneral": limpiar_texto(info),
        "css": limpiar_texto(pd.DataFrame(filas_css[1:], columns=filas_css[0])),
        "colcol": limpiar_texto(read_named_table(wb, "colcol")),
        "contrapartidas": limpiar_texto(contrapartidas),
    }


def cargar_sectores(ruta, hoja_act, hoja_ant):
    """Fuentes de Ficha sectores.xlsx: AOD de ambos cortes (fuente única para las dos fichas, con los
    nombres individuales de los cooperantes), información de los sectores y sector vinculado de ColCol."""
    hojas = pd.read_excel(ruta, sheet_name=["INFO GENERAL", hoja_act, hoja_ant, "COLCOL"])
    return {
        "info": limpiar_texto(hojas["INFO GENERAL"]),
        "aod": limpiar_texto(hojas[hoja_act]),
        "aod_ant": limpiar_texto(hojas[hoja_ant]),
        "colcol": limpiar_texto(hojas["COLCOL"]),
    }


def validar_columnas(tablas):
    """tablas = {nombre: (descripción de la fuente, DataFrame)}. Devuelve los problemas encontrados."""
    problemas = []
    for nombre, (fuente, df) in tablas.items():
        faltan = [c for c in ESQUEMA.get(nombre, []) if buscar_col(df, c) is None]
        if faltan:
            problemas.append(f"{fuente}: faltan las columnas {', '.join(faltan)}")
        elif df.empty:
            problemas.append(f"{fuente}: la tabla está vacía")
    return problemas


# =============================================================
# Reglas de negocio
# =============================================================
def activas_a_fecha(df, fecha_corte):
    """Intervenciones vigentes a la fecha de corte: FECHA INICIAL <= corte <= FECHA FINAL."""
    if df.empty or "FECHA INICIAL" not in df.columns or "FECHA FINAL" not in df.columns:
        return df
    ini = pd.to_datetime(df["FECHA INICIAL"], errors="coerce")
    fin = pd.to_datetime(df["FECHA FINAL"], errors="coerce")
    return df[(ini <= fecha_corte) & (fin >= fecha_corte)]


def contar_movimiento(df_act, df_ant, vigencia_act):
    """Nuevas, continúan, terminadas (fecha final antes de la vigencia actual) y sin reporte."""
    col = "CODIGO INTERVENCION"
    if col not in df_act.columns or col not in df_ant.columns:
        return None
    cod_act = set(df_act[col].dropna().unique())
    cod_ant = set(df_ant[col].dropna().unique())
    salieron = df_ant[df_ant[col].isin(cod_ant - cod_act)]
    fin = pd.to_datetime(salieron["FECHA FINAL"], errors="coerce")
    terminadas = salieron.loc[fin < vigencia_act, col].nunique()
    return {"nuevas": len(cod_act - cod_ant), "continuan": len(cod_act & cod_ant),
            "terminadas": terminadas, "sin_reporte": len(cod_ant - cod_act) - terminadas}


def asignar_sector(df, col, fuente, sectores_norm, codigo_col, equivalencias, registros):
    """Sector de la ficha (normalizado) de cada fila según la configuración.

    1. EQUIVALENCIAS traduce el valor de la fuente al sector de la ficha.
    2. Un valor que no esté en la tabla solo se asigna si coincide exacto con un sector.
    3. Los valores "Por proyecto" y los registros sin sector se completan con REGISTROS (por código).
    """
    if col not in df.columns:
        return pd.Series(None, index=df.index, dtype="object")
    codigos = df[codigo_col] if codigo_col in df.columns else pd.Series(None, index=df.index)

    def uno(valor, codigo):
        v = norm_text(valor)
        destino = equivalencias.get((fuente, v))
        if destino is None:
            destino = v if v in sectores_norm else None
        if destino in (None, "POR PROYECTO"):
            destino = registros.get((fuente, str(codigo).strip()), "")
        return None if destino in DESTINOS_SIN_SECTOR else destino

    return pd.Series([uno(v, c) for v, c in zip(df[col], codigos)], index=df.index, dtype="object")


def unificar_colcol(colcol, colcol_sectores):
    """ColCol del .xlsm (todos los intercambios) con el sector vinculado de Ficha sectores por código."""
    sector = (colcol_sectores[colcol_sectores["SECTOR VINCULADO"].map(norm_text).isin(SIN_DATO) == False]
              .drop_duplicates("CODIGO")
              .assign(CODIGO=lambda d: d["CODIGO"].astype(str).str.strip())
              .set_index("CODIGO")["SECTOR VINCULADO"])
    out = colcol.copy()
    out["SECTOR VINCULADO"] = out["CODIGO"].astype(str).str.strip().map(sector)
    return out


def leer_planes(ruta):
    try:
        return pd.read_excel(ruta).set_index("Divipola")
    except FileNotFoundError:
        return None


def aplicar_planes_sncic(info, planes):
    """Actualiza los campos del plan de trabajo SNCIC de cada departamento con planes_trabajo_sncic.xlsx."""
    if planes is None:
        return info
    info = info.copy()
    c_plan = buscar_col(info, "¿Cuenta con plan de trabajo en el marco del SNCIC?")
    c_cap = buscar_col(info, "Capítulo especial (Plan de trabajo con ciudades dentro del SNCIC)")
    c_cual = buscar_col(info, "Cuál")
    c_anio = buscar_col(info, "Año creación del plan de trabajo")
    for c in (c_plan, c_cap, c_cual, c_anio):
        if c:
            info[c] = info[c].astype(object)
    for i, div in pd.to_numeric(info["Divipola"], errors="coerce").items():
        p = planes.loc[int(div)] if pd.notna(div) and int(div) in planes.index else None
        tiene = p is not None and p["Tiene plan de trabajo"] == "Sí"
        if c_plan:
            info.at[i, c_plan] = "Sí" if tiene else "No"
        if c_cap:
            info.at[i, c_cap] = p["Capítulo especial"] if p is not None else "No"
        if c_cual:
            ciudades = p["Ciudades con capítulo especial"] if p is not None else None
            info.at[i, c_cual] = ciudades if isinstance(ciudades, str) and ciudades else "No aplica"
        if c_anio:
            anio = p["Año de inicio del plan"] if p is not None else None
            if tiene and pd.notna(anio):
                info.at[i, c_anio] = str(int(anio))
            elif not tiene:
                info.at[i, c_anio] = "No aplica"
            # con plan pero sin fecha en la app SNCIC: se conserva el año que ya tenía la ficha
    return info


def registro_publico(fila):
    """Campos de la información general que se muestran (sin Divipola ni columnas internas)."""
    return [(str(k), v) for k, v in fila.items() if norm_text(k) not in CAMPOS_OCULTOS_REGISTRO]


# =============================================================
# Conteos y agregaciones
# =============================================================
def unir(serie):
    """Valores únicos y con dato, separados por punto y coma."""
    vals = {str(v).strip() for v in serie.dropna() if norm_text(v) not in SIN_DATO}
    return "; ".join(sorted(vals))


def n_unicos(df, col):
    return df[col].nunique() if col in df.columns else 0


def n_municipios(df):
    if "MUNICIPIO" not in df.columns:
        return 0
    m = df["MUNICIPIO"].map(norm_text)
    return m[~m.isin(SIN_DATO)].nunique()


def n_departamentos(df):
    if "DEPARTAMENTO" not in df.columns:
        return 0
    return df.loc[df["DEPARTAMENTO"] != AMBITO_NACIONAL, "DEPARTAMENTO"].nunique()


def a_fecha(serie):
    return pd.to_datetime(serie, errors="coerce")


def agregar(df, clave, spec):
    """groupby + agg usando solo las columnas que existen. spec = {salida: (columna, funcion)}."""
    spec = {k: (c, f) for k, (c, f) in spec.items() if c is not None and c in df.columns}
    return df.groupby(clave, sort=False).agg(**spec).reset_index()


def top_by_sum(df, group_col, value_col, n=5):
    if df.empty or group_col not in df.columns or value_col not in df.columns:
        return pd.DataFrame(columns=[group_col, value_col])
    return (df.groupby(group_col, dropna=False)[value_col].sum()
            .sort_values(ascending=False).head(n).reset_index())


def top_comparativo(top_act, df_ant, group_col, col_act, col_ant, value_col="VALOR APORTE (USD)", nombres=None):
    """Top del corte actual con el valor de esos mismos grupos en el corte anterior (0 si no tenían)."""
    if group_col in df_ant.columns and value_col in df_ant.columns:
        ant = df_ant.groupby(group_col)[value_col].sum()
    else:
        ant = pd.Series(dtype=float)
    v_act = top_act[value_col]
    v_ant = top_act[group_col].map(ant).fillna(0)
    return pd.DataFrame({
        group_col: top_act[group_col].map(lambda x: nombres.get(x, x)) if nombres else top_act[group_col],
        col_act: v_act.apply(format_usd),
        col_ant: v_ant.apply(format_usd),
        "Variación": [variacion_txt(a, b) for a, b in zip(v_act, v_ant)],
    })


def proyectos_contrapartida(df):
    """Proyectos únicos: un mismo proyecto aparece una vez por cada departamento donde se ejecuta."""
    claves = [c for c in ("Nombre del proyecto", "Implementador", "Financiador", "AÑO") if c in df.columns]
    return df.drop_duplicates(subset=claves) if claves else df
